from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any


def cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).replace("\u00a0", " ").strip()


def canonical_mli(raw: str) -> str:
    digits = re.sub(r"\D", "", cell(raw))
    if not digits:
        return ""
    return digits.zfill(12)


def header_key(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", cell(h).lower())


def _pick(row: dict[str, str], keys: list[str]) -> str:
    for key in keys:
        if row.get(key):
            return row[key]
    return ""


def _rows_from_xlsx(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [header_key(h) for h in header_row]
    out: list[dict[str, str]] = []
    for i, raw in enumerate(rows_iter, start=2):
        obj: dict[str, str] = {"__row": str(i)}
        any_val = False
        for idx, h in enumerate(headers):
            val = cell(raw[idx] if idx < len(raw) else "")
            if h:
                obj[h] = val
            if val:
                any_val = True
        if any_val:
            out.append(obj)
    return out


def _rows_from_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        header_row = next(reader, None)
        if not header_row:
            return []
        headers = [header_key(h) for h in header_row]
        out: list[dict[str, str]] = []
        for i, raw in enumerate(reader, start=2):
            obj: dict[str, str] = {"__row": str(i)}
            any_val = False
            for idx, h in enumerate(headers):
                val = cell(raw[idx] if idx < len(raw) else "")
                if h:
                    obj[h] = val
                if val:
                    any_val = True
            if any_val:
                out.append(obj)
        return out


def load_table(path: str | Path) -> list[dict[str, str]]:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    suffix = p.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _rows_from_xlsx(p)
    if suffix in {".csv", ".txt"}:
        return _rows_from_csv(p)
    raise ValueError(f"Unsupported file type: {suffix}")


def split_passwords(raw: str) -> list[str]:
    seen: list[str] = []
    for part in cell(raw).split(","):
        pwd = part.strip()
        if pwd and pwd not in seen:
            seen.append(pwd)
    return seen


def merge_password_chain(excel_password: str, fallbacks: list[str] | None) -> list[str]:
    chain: list[str] = []
    primary = cell(excel_password)
    if primary:
        chain.append(primary)
    for pwd in fallbacks or []:
        if pwd and pwd not in chain:
            chain.append(pwd)
    return chain


def parse_credentials(path: str | Path) -> list[dict]:
    seen: dict[str, dict] = {}
    for row in load_table(path):
        mli_id = _pick(row, ["mlimemberid", "mliid", "memberid", "mli"])
        maker_user = _pick(row, ["makeruserid", "makerid", "maker"])
        maker_password = _pick(row, ["makerpassword", "makerpwd", "makerpass"])
        checker_user = _pick(row, ["checkeruserid", "checkerid", "checker"])
        checker_password = _pick(row, ["checkerpassword", "checkerpwd", "checkerpass"])
        if not mli_id or not maker_user:
            continue
        rec = {
            "mliId": mli_id,
            "mliCanonical": canonical_mli(mli_id),
            "makerUser": maker_user,
            "makerPassword": maker_password,
            "makerPasswords": [maker_password] if maker_password else [],
            "checkerUser": checker_user,
            "checkerPassword": checker_password,
            "checkerPasswords": [checker_password] if checker_password else [],
            "sourceRow": row.get("__row", ""),
        }
        if rec["mliCanonical"]:
            seen[rec["mliCanonical"]] = rec
    return list(seen.values())


def parse_claims(path: str | Path) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for row in load_table(path):
        mli_id = _pick(row, ["mliid", "mlimemberid", "memberid", "mli"])
        claim_ref = _pick(row, ["claimrefno", "claimno", "claimreferenceno", "claim"]).upper().replace(" ", "")
        if not mli_id and not claim_ref:
            continue
        key = (canonical_mli(mli_id), claim_ref)
        if claim_ref and key in seen:
            continue
        if claim_ref:
            seen.add(key)
        out.append(
            {
                "id": f"{canonical_mli(mli_id)}:{claim_ref}:{row.get('__row','')}",
                "mliId": mli_id,
                "mliCanonical": canonical_mli(mli_id),
                "claimRef": claim_ref,
                "state": _pick(row, ["state", "branchstate"]),
                "legalWaiver": _pick(row, ["legalwaiver", "waiver"]),
                "legalForum": _pick(row, ["legalforum", "forum", "legalproceedings"]),
                "attachment": _pick(row, ["attachment", "legalattachment", "legalfile", "attachmentpath"]),
                "comment": _pick(row, ["comment", "remarks", "mlcomment"]),
                "urn": _pick(row, ["urn", "urnno", "urnnumber"]),
                "sample": "",
                "sourceRow": row.get("__row", ""),
            }
        )
    return out
